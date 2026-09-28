from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import patch

from fastapi import HTTPException
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import Session, sessionmaker
from sqlalchemy.pool import StaticPool

from app.api.routes.conciliaciones import (
    approve_reconciliation,
    export_results,
    reject_reconciliation,
)
from app.api.routes.ejecuciones import read_ejecucion
from app.database.base import Base
from app.models import (
    Cliente,
    EjecucionProceso,
    Proceso,
    ResultadoConciliacion,
    Usuario,
)
from app.schemas.resultado_revision import RechazarEjecucionRequest


class ConciliacionClosureTests(unittest.TestCase):
    def setUp(self) -> None:
        self.engine = create_engine(
            "sqlite://",
            connect_args={"check_same_thread": False},
            poolclass=StaticPool,
        )
        Base.metadata.create_all(self.engine)
        self.session_factory = sessionmaker(
            bind=self.engine,
            class_=Session,
            expire_on_commit=False,
        )
        self.temp_directory = TemporaryDirectory()
        self._seed_database()

    def tearDown(self) -> None:
        self.engine.dispose()
        self.temp_directory.cleanup()

    def _seed_database(self) -> None:
        with self.session_factory() as db:
            owner_client = Cliente(nombre="Cliente cierre A")
            other_client = Cliente(nombre="Cliente cierre B")
            db.add_all([owner_client, other_client])
            db.flush()
            self.owner = Usuario(
                cliente_id=owner_client.id,
                nombre="Admin A",
                email="admin-a@example.test",
                password_hash="not-used",
                rol="ADMIN",
                estado="ACTIVO",
            )
            self.other = Usuario(
                cliente_id=other_client.id,
                nombre="Admin B",
                email="admin-b@example.test",
                password_hash="not-used",
                rol="ADMIN",
                estado="ACTIVO",
            )
            process = Proceso(
                cliente_id=owner_client.id,
                nombre="Conciliación cierre",
                tipo="CONCILIACION_EXCEL",
            )
            db.add_all([self.owner, self.other, process])
            db.flush()
            execution = EjecucionProceso(
                proceso_id=process.id,
                usuario_id=self.owner.id,
                estado="REQUIERE_REVISION",
                resumen_json={"conciliacion_resumen": {"total_resultados": 1}},
            )
            db.add(execution)
            db.flush()
            result = ResultadoConciliacion(
                ejecucion_id=execution.id,
                clave_referencia="FAC-1",
                estado_resultado="DIFERENCIA_IMPORTE",
                diferencia_importe=20,
                requiere_revision=False,
                observacion="Validado",
            )
            db.add(result)
            db.commit()
            self.execution_id = execution.id
            self.result_id = result.id

    def _set_state(self, state: str) -> None:
        with self.session_factory() as db:
            execution = db.get(EjecucionProceso, self.execution_id)
            execution.estado = state
            db.commit()

    def test_owner_can_approve_when_review_is_complete(self) -> None:
        with self.session_factory() as db:
            summary = approve_reconciliation(self.execution_id, db, self.owner)

        self.assertEqual(summary["estado_ejecucion"], "APROBADO")
        with self.session_factory() as db:
            self.assertEqual(
                db.get(EjecucionProceso, self.execution_id).estado,
                "APROBADO",
            )

    def test_approval_rejects_pending_reviews_and_terminal_states(self) -> None:
        with self.session_factory() as db:
            result = db.get(ResultadoConciliacion, self.result_id)
            result.requiere_revision = True
            db.commit()
            with self.assertRaises(HTTPException) as pending_context:
                approve_reconciliation(self.execution_id, db, self.owner)
        self.assertEqual(pending_context.exception.status_code, 400)

        self._set_state("RECHAZADO")
        with self.session_factory() as db:
            with self.assertRaises(HTTPException) as terminal_context:
                approve_reconciliation(self.execution_id, db, self.owner)
        self.assertEqual(terminal_context.exception.status_code, 409)

    def test_owner_can_reject_and_reason_is_persisted(self) -> None:
        with self.session_factory() as db:
            summary = reject_reconciliation(
                self.execution_id,
                RechazarEjecucionRequest(motivo="Documentación inconsistente"),
                db,
                self.owner,
            )

        self.assertEqual(summary["estado_ejecucion"], "RECHAZADO")
        with self.session_factory() as db:
            execution = db.get(EjecucionProceso, self.execution_id)
            self.assertEqual(execution.error_message, "Documentación inconsistente")
            self.assertEqual(
                execution.resumen_json["rechazo"],
                {
                    "motivo": "Documentación inconsistente",
                    "usuario_id": self.owner.id,
                },
            )

    def test_terminal_execution_cannot_be_rejected_again(self) -> None:
        self._set_state("APROBADO")
        with self.session_factory() as db:
            with self.assertRaises(HTTPException) as context:
                reject_reconciliation(
                    self.execution_id,
                    RechazarEjecucionRequest(motivo="Tardío"),
                    db,
                    self.owner,
                )
        self.assertEqual(context.exception.status_code, 409)

    def test_cross_tenant_closure_export_and_status_are_concealed(self) -> None:
        operations = (
            lambda db: approve_reconciliation(self.execution_id, db, self.other),
            lambda db: reject_reconciliation(
                self.execution_id,
                RechazarEjecucionRequest(motivo="Ajeno"),
                db,
                self.other,
            ),
            lambda db: export_results(self.execution_id, db, self.other),
            lambda db: read_ejecucion(self.execution_id, db, self.other),
        )
        for operation in operations:
            with self.subTest(operation=operation):
                with self.session_factory() as db:
                    with self.assertRaises(HTTPException) as context:
                        operation(db)
                self.assertEqual(context.exception.status_code, 404)

    def test_approved_owner_can_export_xlsx_with_preserved_headers(self) -> None:
        self._set_state("APROBADO")
        output_root = Path(self.temp_directory.name) / "processed"
        with self.session_factory() as db, patch(
            "app.services.conciliacion_export_service.PROCESSED_STORAGE_ROOT",
            output_root,
        ):
            response = export_results(self.execution_id, db, self.owner)

        output_path = Path(response.path)
        self.assertTrue(output_path.is_file())
        self.assertEqual(
            response.media_type,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn(
            f"conciliacion_ejecucion_{self.execution_id}.xlsx",
            response.headers["content-disposition"],
        )
        workbook = load_workbook(output_path, read_only=True)
        self.assertIn("Resumen", workbook.sheetnames)
        self.assertIn("Todos", workbook.sheetnames)
        workbook.close()

    def test_export_rejects_unapproved_and_stale_execution(self) -> None:
        for state in ("REQUIERE_REVISION", "CARGADO", "RECHAZADO"):
            with self.subTest(state=state):
                self._set_state(state)
                with self.session_factory() as db:
                    with self.assertRaises(HTTPException) as context:
                        export_results(self.execution_id, db, self.owner)
                self.assertEqual(context.exception.status_code, 409)


if __name__ == "__main__":
    unittest.main()
